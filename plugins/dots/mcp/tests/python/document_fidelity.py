#!/usr/bin/env python3
"""Fidelity/security-input tests on generated fixtures. This does NOT qualify OS parser containment.
The production broker never routes untrusted work through this direct test runner.
"""
import sys,os,json,subprocess,unittest,tempfile,zipfile,hashlib,io,warnings
from pathlib import Path
from xml.etree import ElementTree as E
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from PIL import Image
from pypdf import PdfReader,PdfWriter
from reportlab.pdfgen.canvas import Canvas
ROOT=Path(__file__).resolve().parents[2]
WORKER=ROOT/'packages/capabilities/src/documents/worker.py'
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def parts(p):
 with zipfile.ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
def save_parts(p,items):
 with zipfile.ZipFile(p,'w',compression=zipfile.ZIP_DEFLATED) as z:
  for name,b in items.items():z.writestr(name,b)
def make_fixtures(directory):
 d=Path(directory);d.mkdir(parents=True,exist_ok=True)
 im=Image.new('RGB',(96,64),'white')
 for x in range(96):
  for y in range(64):im.putpixel((x,y),(x*2,y*3,(x+y)%255))
 im.save(d/'fixture.png');im.save(d/'fixture.webp');im.save(d/'fixture.jpg')
 two=Image.new('RGB',(96,64),'black');im.save(d/'animated.gif',save_all=True,append_images=[two],duration=100,loop=0)
 doc=Document();doc.sections[0].header.paragraphs[0].text='Unchanged header';doc.sections[0].footer.paragraphs[0].text='Unchanged footer'
 doc.add_heading('Fixture report',0);p=doc.add_paragraph();p.add_run('Target ');run=p.add_run('value');run.bold=True;p.add_run(' remains styled')
 doc.add_paragraph('Unicode: café — résumé');table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Metric';table.cell(0,1).text='Amount';table.cell(1,0).text='Revenue';table.cell(1,1).text='100'
 doc.add_picture(str(d/'fixture.png'),width=Inches(1));doc.save(d/'rich.docx')
 wb=Workbook();ws=wb.active;ws.title='Inputs';ws['A1']='Revenue';ws['A1'].font=Font(bold=True);ws['A1'].fill=PatternFill('solid',fgColor='DDDDDD');ws['B1']=100;ws['C1']='=B1*2';ws['B1'].number_format='#,##0.00';keep=wb.create_sheet('Untouched');keep['A1']='keep exactly';keep['B2']='=Inputs!B1';wb.save(d/'rich.xlsx')
 z=parts(d/'rich.xlsx');root=E.fromstring(z['xl/worksheets/sheet1.xml']);cell=next(x for x in root.iter('{'+S+'}c') if x.get('r')=='C1');cell.find('{'+S+'}v').text='200';z['xl/worksheets/sheet1.xml']=E.tostring(root);save_parts(d/'rich.xlsx',z)
 z['xl/vbaProject.bin']=b'INERT VBA FIXTURE - NOT EXECUTABLE';save_parts(d/'inert.xlsm',z)
 c=Canvas(str(d/'rich.pdf'));c.drawString(70,740,'Fixture page one');c.drawImage(str(d/'fixture.png'),70,600,96,64);c.showPage();c.drawString(70,740,'Fixture page two');c.showPage();c.drawString(70,740,'Fixture page three');c.save()
 writer=PdfWriter();writer.append(str(d/'rich.pdf'));writer.encrypt('fixture-only-password');writer.write(str(d/'encrypted.pdf'))
 (d/'corrupt.docx').write_bytes(b'not a ZIP');(d/'legacy.xls').write_bytes(b'fixture-not-a-supported-binary-workbook')
 with zipfile.ZipFile(d/'traversal.docx','w') as z:z.writestr('../escape','x')
 with zipfile.ZipFile(d/'bomb.docx','w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('bomb.bin',b'0'*1048576)
 with zipfile.ZipFile(d/'entity.docx','w') as z:z.writestr('word/document.xml',b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')
 manifest={p.name:{'bytes':p.stat().st_size,'sha256':digest(p)} for p in d.iterdir() if p.is_file() and p.name!='manifest.json'};(d/'manifest.json').write_text(json.dumps({'generated':True,'containsUserData':False,'fixtures':manifest},indent=2))
 return d
class Fidelity(unittest.TestCase):
 @classmethod
 def setUpClass(cls):cls.tmp=tempfile.TemporaryDirectory(prefix='portal-document-tests-');cls.d=make_fixtures(cls.tmp.name)
 @classmethod
 def tearDownClass(cls):cls.tmp.cleanup()
 def run_worker(self,family,mode,input=None,success=True,**args):
  out=self.d/('out-'+str(len(list(self.d.iterdir())))+('.png' if mode=='render' else '.bin'));payload={'family':family,'mode':mode,'output':str(out),**args}
  if input:payload['input']=str(self.d/input)
  r=subprocess.run([sys.executable,'-I',str(WORKER)],input=json.dumps(payload).encode(),capture_output=True,timeout=30,cwd=self.d,env={'PATH':'/usr/bin:/bin','HOME':str(self.d),'LANG':'C.UTF-8'})
  result=json.loads(r.stdout)
  if success:self.assertEqual(r.returncode,0,(r.stderr.decode(),result));self.assertNotIn('error',result)
  else:self.assertNotEqual(r.returncode,0);self.assertIn('error',result)
  return result,out
 def test_D01_docx_outline_tables_and_actual_text(self):
  result,_=self.run_worker('docx','inspect','rich.docx');self.assertTrue(any(x['kind']=='table' for x in result['elements']));self.assertTrue(any('Target value' in x.get('text','') for x in result['elements']))
 def test_D02_docx_targeted_edit_preserves_all_other_zip_parts(self):
  before=parts(self.d/'rich.docx');result,out=self.run_worker('docx','edit','rich.docx',edits=[{'locator':'p:1','oldText':'value','newText':'improved','expectedReplacements':1}]);after=parts(out)
  self.assertEqual(set(before),set(after));self.assertEqual(result['changedParts'],['word/document.xml'])
  for name in before:
   if name!='word/document.xml':self.assertEqual(before[name],after[name],name)
  parsed=Document(out);self.assertIn('Target improved',parsed.paragraphs[1].text);self.assertTrue(parsed.paragraphs[1].runs[1].bold)
  # Prefixes named by mc:Ignorable must remain declared, not merely visually plausible.
  from lxml import etree
  root=etree.fromstring(after['word/document.xml']);ign=root.attrib.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable','')
  for prefix in ign.split():self.assertIn(prefix,root.nsmap)
 def test_D03_docx_complex_cross_run_edit_is_rejected(self):
  self.run_worker('docx','edit','rich.docx',success=False,edits=[{'locator':'p:1','oldText':'Target value','newText':'new','expectedReplacements':1}])
 def test_D04_docx_table_cell_edit_and_creation(self):
  _,out=self.run_worker('docx','edit','rich.docx',edits=[{'locator':'t:0:r:1:c:1','oldText':'100','newText':'120','expectedReplacements':1}]);self.assertEqual(Document(out).tables[0].cell(1,1).text,'120')
  _,created=self.run_worker('docx','create',markdown='# Title\nBody café',tables=[[['a','b'],['c','d']]]);self.assertEqual(Document(created).paragraphs[0].text,'Title');self.assertEqual(Document(created).tables[0].cell(1,1).text,'d')
 def test_D05_sheet_formula_and_cached_value_are_distinct(self):
  result,_=self.run_worker('sheet','read','rich.xlsx',sheet='Inputs',range='A1:C1');formula=next(c for c in result['cells'] if c['cell']=='C1');self.assertEqual(formula['formula'],'B1*2');self.assertEqual(formula['cachedValue'],200);self.assertEqual(formula['cacheStatus'],'not-verified');self.assertTrue(result['recalculationRequired'])
 def test_D06_sheet_targeted_edit_preserves_untouched_parts(self):
  before=parts(self.d/'rich.xlsx');result,out=self.run_worker('sheet','edit','rich.xlsx',sheet='Inputs',cells=[{'cell':'B1','value':125}]);after=parts(out)
  for name in before:
   if name not in ['xl/workbook.xml','xl/worksheets/sheet1.xml']:self.assertEqual(before[name],after[name],name)
  book=load_workbook(io.BytesIO(out.read_bytes()),data_only=False);self.assertEqual(book['Inputs']['B1'].value,125);self.assertEqual(book['Inputs']['C1'].value,'=B1*2');self.assertEqual(book['Untouched']['B2'].value,'=Inputs!B1');self.assertTrue(result['recalculationRequired'])
 def test_D07_sheet_formula_write_has_no_fabricated_cached_result(self):
  _,out=self.run_worker('sheet','edit','rich.xlsx',sheet='Inputs',cells=[{'cell':'C1','formula':'=B1*3','format':{'bold':True,'numberFormat':'0.00'}}]);book=load_workbook(io.BytesIO(out.read_bytes()),data_only=False);self.assertEqual(book['Inputs']['C1'].value,'=B1*3');self.assertTrue(book['Inputs']['C1'].font.bold);self.assertIsNone(load_workbook(io.BytesIO(out.read_bytes()),data_only=True)['Inputs']['C1'].value)
 def test_D08_macro_payload_is_read_only_never_executed_or_discarded(self):
  before=digest(self.d/'inert.xlsm');result,_=self.run_worker('sheet','inspect','inert.xlsm');self.assertTrue(result['macrosPresent']);self.assertFalse(result['editingAvailable']);self.assertFalse(result['macroExecution']);self.run_worker('sheet','edit','inert.xlsm',success=False,sheet='Inputs',cells=[{'cell':'B1','value':999}]);self.assertEqual(digest(self.d/'inert.xlsm'),before)
 def test_D09_sheet_creation_invalid_ranges_and_duplicate_names(self):
  _,out=self.run_worker('sheet','create',sheets=[{'name':'One','cells':[{'cell':'A1','value':'001'},{'cell':'B1','value':True},{'cell':'C1','formula':'=1+1'}]}]);book=load_workbook(io.BytesIO(out.read_bytes()));self.assertEqual(book['One']['A1'].value,'001');self.assertEqual(book['One']['B1'].value,True)
  self.run_worker('sheet','read','rich.xlsx',success=False,sheet='Inputs',range='A1:XFD1048576');self.run_worker('sheet','create',success=False,sheets=[{'name':'same'},{'name':'SAME'}])
 def test_D10_pdf_inspect_extract_and_real_page_render(self):
  result,_=self.run_worker('pdf','inspect','rich.pdf');self.assertEqual(result['pages'],3);read,_=self.run_worker('pdf','read','rich.pdf',startPage=2,maxPages=1);self.assertIn('Fixture page two',read['pages'][0]['text']);render,out=self.run_worker('pdf','render','rich.pdf',page=1,maxEdge=800)
  with Image.open(out) as image:self.assertEqual(image.format,'PNG');self.assertLessEqual(max(image.size),800);self.assertGreater(len(image.convert('RGB').getcolors(1000000)),2)
  self.assertEqual(render['page'],1)
 def test_D11_pdf_page_reordering_keeps_original_and_text_order(self):
  original=digest(self.d/'rich.pdf');result,out=self.run_worker('pdf','edit','rich.pdf',pages=[{'page':3},{'page':1}]);reader=PdfReader(out);self.assertEqual(len(reader.pages),2);self.assertIn('three',reader.pages[0].extract_text());self.assertIn('one',reader.pages[1].extract_text());self.assertEqual(digest(self.d/'rich.pdf'),original)
 def test_D12_pdf_creation_and_encrypted_input_rejection(self):
  _,out=self.run_worker('pdf','create',markdown='# Created report\nActual content');self.assertIn('Actual content',PdfReader(out).pages[0].extract_text());self.run_worker('pdf','inspect','encrypted.pdf',success=False)
 def test_D13_image_crop_resize_formats_and_first_animation_frame(self):
  for name in ['fixture.png','fixture.jpg','fixture.webp','animated.gif']:
   result,_=self.run_worker('image','inspect',name);self.assertEqual(result['width'],96)
  result,out=self.run_worker('image','view','fixture.png',crop={'x':10,'y':10,'width':60,'height':40},maxEdge=30)
  with Image.open(out) as image:self.assertEqual(image.size,(30,20))
  result,_=self.run_worker('image','view','animated.gif',maxEdge=32);self.assertEqual(result['animation'],'first-frame');self.run_worker('image','view','fixture.png',success=False,crop={'x':90,'y':60,'width':60,'height':40})
 def test_D14_archive_corruption_traversal_bombs_and_entities_fail(self):
  for name in ['corrupt.docx','traversal.docx','bomb.docx','entity.docx']:self.run_worker('docx','inspect',name,success=False)
 def test_D15_duplicate_zip_parts_and_utf16_entities_are_rejected(self):
  with warnings.catch_warnings():
   warnings.simplefilter('ignore')
   with zipfile.ZipFile(self.d/'duplicate.docx','w') as z:z.writestr('word/document.xml','<x/>');z.writestr('word/document.xml','<y/>')
  self.run_worker('docx','inspect','duplicate.docx',success=False)
  with zipfile.ZipFile(self.d/'utf16.docx','w') as z:z.writestr('word/document.xml','<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY a "bad">]><x>&a;</x>'.encode('utf-16'))
  self.run_worker('docx','inspect','utf16.docx',success=False)
if __name__=='__main__':
 if len(sys.argv)>1 and sys.argv[1]=='--generate':print(make_fixtures(sys.argv[2]));sys.exit(0)
 unittest.main(verbosity=2)
